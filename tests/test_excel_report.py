# -*- coding: utf-8 -*-
"""Excel 报告: 懒建 sheet / 两段式步骤名 / 汇总统计 / 截图嵌入。"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from core.excel_report import ExcelReport
from openpyxl import load_workbook


def test_two_phase_desc_and_stats(tmp_path):
    p = tmp_path / "r.xlsx"
    r = ExcelReport(path=p)
    r.set_step_desc("全局清扫", "点击开始清扫")
    r.add_result("全局清扫", True)
    r.set_step_desc("全局清扫", "检查清扫中")
    r.add_result("全局清扫", False, error="超时")
    r.save()
    assert p.exists()

    from openpyxl import load_workbook
    wb = load_workbook(p)
    ws = wb["全局清扫"]
    assert ws.cell(row=2, column=1).value == "点击开始清扫"
    assert ws.cell(row=2, column=2).value == "PASS"
    assert ws.cell(row=3, column=1).value == "检查清扫中"
    assert ws.cell(row=3, column=2).value == "FAIL"
    assert ws.cell(row=3, column=3).value == "超时"


def test_summary_sheet_has_pass_rate(tmp_path):
    p = tmp_path / "r.xlsx"
    r = ExcelReport(path=p)
    r.add_result("用例A", True)
    r.add_result("用例B", False)
    r.save()
    from openpyxl import load_workbook
    wb = load_workbook(p)
    ws = wb["汇总"]
    texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert any("通过率" in t for t in texts)
    assert any("设备SN" in t for t in texts)


def test_screenshot_embedded(tmp_path):
    from PIL import Image
    img = tmp_path / "shot.png"
    Image.new("RGB", (40, 30), "red").save(img)
    p = tmp_path / "r2.xlsx"
    r = ExcelReport(path=p)
    r.set_step_desc("用例A", "截图步骤")
    r.add_result("用例A", True, screenshot=str(img))
    r.save()
    from openpyxl import load_workbook
    wb = load_workbook(p)
    assert len(wb["用例A"]._images) == 1


def test_device_info_filled(tmp_path):
    p = tmp_path / "r3.xlsx"
    r = ExcelReport(path=p)
    r.set_device_info(sn="SN12345", app_version="1.8.45.24")
    r.add_result("用例A", True)
    r.save()
    from openpyxl import load_workbook
    wb = load_workbook(p)
    texts = [str(c.value) for row in wb["汇总"].iter_rows() for c in row if c.value]
    assert any("SN12345" in t for t in texts)


def test_case_name_sanitized(tmp_path):
    """用例名含 / \\ 时必须清洗,否则 openpyxl 建 sheet 报错"""
    p = tmp_path / "r4.xlsx"
    r = ExcelReport(path=p)
    r.add_result("清扫/记录", True)
    r.save()
    from openpyxl import load_workbook
    wb = load_workbook(p)
    assert "清扫_记录" in wb.sheetnames


def test_sheet_name_sanitizes_all_openpyxl_forbidden_chars(tmp_path):
    """★ sheet 名要清掉 openpyxl 全部非法字符([ ] : * ? / \\), 不能让整轮挂掉。

    回归守护: 原实现只替换了 / 和 \\ —— 组名写「温度: 设置」时 create_sheet 直接抛
    `ValueError: Invalid character : found in sheet title`, 冒到 RunWorker.run 的
    兜底 except → **整轮**以"执行异常"中断, 设备侧已跑的部分白跑。
    组名(module)是 GUI 里可自由编辑的文本, 很容易踩到。
    """
    from openpyxl import load_workbook

    from core.excel_report import ExcelReport
    p = tmp_path / "r.xlsx"
    r = ExcelReport(path=p)
    names = ["温度: 设置", "模式[夜间]", "速度?快", "a*b", "a\\b", "a/b", "全局清扫"]
    for n in names:
        r.add_result(n, True)          # 关键: 这里不得抛异常
    r.save()
    wb = load_workbook(p)
    # 「汇总」是报告自带的汇总页; a\b 与 a/b 清洗后同名, openpyxl 会自己加后缀
    case_sheets = [s for s in wb.sheetnames if s != "汇总"]
    assert len(case_sheets) == len(names)
    for sn in wb.sheetnames:
        assert not (set(sn) & set('[]:*?/\\')), f"sheet 名仍含非法字符: {sn!r}"


def test_summary_case_name_links_to_its_detail_sheet(tmp_path):
    """★ 汇总页的用例名称要能**点进**那一次执行的详情页(用户要求 2026-10-12)。

    同一个用例跑多轮时, 汇总页每行是一个"第几次执行", 链接必须指向**那一行对应的**
    详情页(而不是笼统的用例名 —— 那样只会跳到第一轮)。
    """
    from openpyxl import load_workbook
    from core.excel_report import ExcelReport

    rep = ExcelReport(path=str(tmp_path / "r.xlsx"))
    rep.add_result("清扫-第1次", True, "", "")
    rep.add_result("清扫-第2次", False, "炸了", "")
    rep.save()

    wb = load_workbook(tmp_path / "r.xlsx")
    ws = wb["汇总"]
    # 汇总数据从第 6 行开始
    for row, sheet, result in ((6, "清扫-第1次", "PASS"), (7, "清扫-第2次", "FAIL")):
        cell = ws.cell(row=row, column=2)
        assert cell.value == sheet, cell.value
        assert cell.hyperlink is not None, f"第{row}行的用例名没有超链接"
        assert cell.hyperlink.location == f"'{sheet}'!A1", cell.hyperlink.location
        assert cell.hyperlink.target is None, "内部跳转不该有外部 target"
        assert cell.font.underline == "single" and cell.font.color.rgb.endswith("0563C1"), \
            "链接该是蓝色下划线"
        assert ws.cell(row=row, column=4).value == result
    # 两张详情页都在
    assert "清扫-第1次" in wb.sheetnames and "清扫-第2次" in wb.sheetnames


def test_summary_link_escapes_single_quote_in_sheet_name(tmp_path):
    """表名里有单引号时要按 Excel 规则写成两个, 否则链接会指错/报错"""
    from openpyxl import load_workbook
    from core.excel_report import ExcelReport

    rep = ExcelReport(path=str(tmp_path / "r.xlsx"))
    rep.add_result("张三's 用例", True, "", "")
    rep.save()
    wb = load_workbook(tmp_path / "r.xlsx")
    cell = wb["汇总"].cell(row=6, column=2)
    assert cell.hyperlink.location == "'张三''s 用例'!A1", cell.hyperlink.location


def test_summary_shows_case_end_time(tmp_path):
    """★ 汇总页要显示每个用例(每次执行)的**结束时间**(用户要求 2026-10-12)。

    时间取"该次执行最后一条结果写入的时刻" —— 报告是在每步跑完时写的, 秒级足够。
    """
    from openpyxl import load_workbook
    from core.excel_report import ExcelReport

    rep = ExcelReport(path=str(tmp_path / "r.xlsx"))
    rep.add_result("清扫-第1次", True, "", "")
    rep.add_result("清扫-第1次", True, "", "")
    rep.add_result("清扫-第2次", False, "炸了", "")
    rep.save()

    wb = load_workbook(tmp_path / "r.xlsx")
    ws = wb["汇总"]
    assert ws.cell(row=5, column=8).value == "结束时间", "表头缺「结束时间」"
    t1 = ws.cell(row=6, column=8).value
    t2 = ws.cell(row=7, column=8).value
    import re as _re
    assert _re.fullmatch(r"\d{2}-\d{2} \d{2}:\d{2}:\d{2}", str(t1)), t1
    assert _re.fullmatch(r"\d{2}-\d{2} \d{2}:\d{2}:\d{2}", str(t2)), t2
    # 超链接仍在(同一列的用例名)
    assert ws.cell(row=6, column=2).hyperlink is not None


def test_each_case_execution_has_its_own_end_time(tmp_path):
    """★ 每条用例(每次执行)的结束时间必须是**各自的**, 不能全变成一个值。

    用户 2026-10-12 特意提醒: "注意是记录每条用例的单独结束时间, 不要又变成同一个结束时间了"
    —— 报告里每个"第几次执行"是**独立的 sheet**(case_sheet_name 生成), 时间在建表时就分开
    存在各自的 case 记录里; 这条用"三次执行之间真的隔了时间"来钉住它, 而不是只看格式合法。
    """
    import re as _re
    import time as _t

    from openpyxl import load_workbook
    from core.excel_report import ExcelReport

    rep = ExcelReport(path=str(tmp_path / "r.xlsx"))
    rep.add_result("清扫-第1次", True, "", "")
    _t.sleep(1.05)                       # 让秒数真的跨过去
    rep.add_result("清扫-第2次", True, "", "")
    _t.sleep(1.05)
    rep.add_result("清扫-第3次", False, "炸了", "")
    rep.save()

    ws = load_workbook(tmp_path / "r.xlsx")["汇总"]
    times = [ws.cell(row=r, column=8).value for r in (6, 7, 8)]
    for t in times:
        assert _re.fullmatch(r"\d{2}-\d{2} \d{2}:\d{2}:\d{2}", str(t)), times
    assert len(set(times)) == 3, f"三次执行的时间不该相同: {times}"


def test_same_case_twice_shares_no_state(tmp_path):
    """同一个用例的多次执行**不共享状态**: 各自的步骤只进自己的详情页, 统计各自算"""
    from openpyxl import load_workbook
    from core.excel_report import ExcelReport

    rep = ExcelReport(path=str(tmp_path / "r.xlsx"))
    rep.set_step_desc("甲-第1次", "第一次的步骤")
    rep.add_result("甲-第1次", True, "", "")
    rep.set_step_desc("甲-第2次", "第二次的步骤")
    rep.add_result("甲-第2次", True, "", "")
    rep.add_result("甲-第2次", True, "", "")          # 第二次有两条
    rep.save()

    wb = load_workbook(tmp_path / "r.xlsx")
    first = [c.value for row in wb["甲-第1次"].iter_rows() for c in row if c.value]
    second = [c.value for row in wb["甲-第2次"].iter_rows() for c in row if c.value]
    assert "第一次的步骤" in first and "第二次的步骤" not in first, first
    assert "第二次的步骤" in second and "第一次的步骤" not in second, second
    assert first.count("PASS") == 1, first            # 步骤数各自独立
    assert second.count("PASS") == 2, second
    ws = wb["汇总"]                                    # 汇总页逐行对应一次执行
    assert (ws.cell(row=6, column=5).value, ws.cell(row=6, column=6).value) == (1, 1)
    assert (ws.cell(row=7, column=5).value, ws.cell(row=7, column=6).value) == (2, 2)


def test_case_priority_written_to_summary(tmp_path):
    """★ 汇总页"优先级"列写真实值(用户 2026-10-15: "这里设置的优先级也没有在
    测试报告的汇总页写出来") —— 界面/用例 YAML 里设置的 priority 此前恒为空。"""
    from core.excel_report import ExcelReport
    rep = ExcelReport(path=str(tmp_path / "p.xlsx"))
    rep.set_case_priority("冒烟", "P0")
    rep.add_result("冒烟", True, "", "")
    rep.add_result("别的用例", True, "", "")      # 没设优先级 → 留空
    rep.save()
    wb = load_workbook(tmp_path / "p.xlsx")
    ws = wb["汇总"]
    rows = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3).value
            for r in range(6, ws.max_row + 1) if ws.cell(row=r, column=2).value}
    assert rows.get("冒烟") == "P0", rows
    assert rows.get("别的用例") in ("", None), rows
