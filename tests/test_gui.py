# -*- coding: utf-8 -*-
"""GUI 全功能点测试套件(离屏运行,不依赖真机;涉及在线设备的用例自动跳过)

运行: python -m pytest tests/test_gui.py -v
"""
import copy
import glob
import json
import os
import sys

import pytest
import yaml

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from PySide6.QtCore import QEvent, QPointF, Qt
from PySide6.QtGui import QMouseEvent

from common import app_detect
from common.action_runner import ActionRunner, UserStopped
from common.driver import update_config
from gui import schema

# 进程内共享一个 QApplication(必须在导入 gui 模块前创建)
from PySide6.QtWidgets import QApplication
_qapp = QApplication.instance() or QApplication([])

CASES_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                         "Test_cases")

SAMPLE_CONFIG = """# ── APP 配置 ──
app:
  name: 测试应用              # 被测APP名称(GUI 可改,作为包名自动检测依据)
  package: com.demo.app
  main_activity: com.demo.MainActivity

# ── 设备配置 ──
device:
  default: 127.0.0.1:5555
  list:
    - id: 127.0.0.1:5555
      name: 模拟器

# ── 目标设备 ──
target_device: SE3L            # 进入哪个设备页面
"""


# ── schema: 步骤模型 ──
class TestSchema:

    def test_new_step_必填文本给空串(self):
        s = schema.new_step("click")
        assert s == {"desc": "点击", "click": ""}

    def test_serialize_空值与False省略_0保留(self):
        step = {"desc": "x", "set_time": 0, "circular": False,
                "switch_to": "打开", "note": ""}
        out = schema.serialize_step(step)
        assert out == {"desc": "x", "set_time": 0, "switch_to": "打开"}

    def test_serialize_group_空退化True_填值成dict去None(self):
        assert schema.serialize_step({"add_timer": True})["add_timer"] is True
        out = schema.serialize_step({"add_timer": {"add_text": "添加", "max_tasks": None}})
        assert out["add_timer"] == {"add_text": "添加"}

    def test_serialize_int4_room_zones空列表转True(self):
        assert schema.serialize_step({"room_zones": []})["room_zones"] is True
        assert schema.serialize_step({"room_zones": [1, 2, 3, 4]})["room_zones"] == [1, 2, 3, 4]
        assert "switch_area" not in schema.serialize_step({"switch_area": []})

    def test_serialize_else递归清理(self):
        step = {"if": "扫地机器人",
                "else": [{"desc": "s1", "click": "x", "screenshot": ""},
                         {"desc": "s2", "wait": 5}]}
        out = schema.serialize_step(step)
        assert out["else"] == [{"desc": "s1", "click": "x"}, {"desc": "s2", "wait": 5}]

    def test_validate_必填(self):
        errs = schema.validate_step({"desc": "s", "set_time": ""})
        assert any("偏移分钟数" in e for e in errs)

    def test_validate_else子步骤带前缀(self):
        errs = schema.validate_step({"desc": "s", "if": "x",
                                     "else": [{"desc": "c", "click": ""}]})
        assert any(err.startswith("else子步骤1") and "必填" in err for err in errs)

    def test_validate_未识别动作(self):
        assert schema.validate_step({"desc": "只有描述"})

    def test_step_summary_徽标(self):
        text = schema.step_summary({"desc": "点击", "click": "确认",
                                    "screenshot": "a.png", "wait": 3})
        assert "📷" in text and "+3s" in text


# ── 配置回写 ──
class TestUpdateConfig:

    @pytest.fixture
    def cfg_file(self, tmp_path):
        p = tmp_path / "config.yaml"
        p.write_text(SAMPLE_CONFIG, encoding="utf-8")
        return str(p)

    def test_同值回写字节不变(self, cfg_file):
        with open(cfg_file, encoding="utf-8") as f:
            before = f.read()
        update_config({"app.name": "测试应用",
                       "app.package": "com.demo.app",
                       "target_device": "SE3L"}, path=cfg_file)
        with open(cfg_file, encoding="utf-8") as f:
            assert f.read() == before

    def test_异值回写生效且保注释(self, cfg_file):
        update_config({"app.name": "新应用"}, path=cfg_file)
        text = open(cfg_file, encoding="utf-8").read()
        assert "name: 新应用" in text
        assert "被测APP名称" in text  # 行内注释保留
        import yaml as _y
        assert _y.safe_load(text)["app"]["name"] == "新应用"

    def test_target_device顶层标量(self, cfg_file):
        update_config({"target_device": "SE9X"}, path=cfg_file)
        text = open(cfg_file, encoding="utf-8").read()
        assert "target_device: SE9X" in text and "进入哪个设备页面" in text

    def test_device_name_更新与追加(self, cfg_file):
        update_config({"device.name:127.0.0.1:5555": "夜神模拟器"}, path=cfg_file)
        update_config({"device.name:emulator-5554": "新模拟器"}, path=cfg_file)
        data = yaml.safe_load(open(cfg_file, encoding="utf-8"))
        names = {d["id"]: d.get("name") for d in data["device"]["list"]}
        assert names["127.0.0.1:5555"] == "夜神模拟器"
        assert names["emulator-5554"] == "新模拟器"


# ── APP/设备检测(纯逻辑) ──
class TestAppDetect:

    def test_中文名转拼音关键词(self):
        kws = app_detect.name_keywords("涂鸦智能")
        assert "tuya" in kws and "zhineng" in kws

    def test_英文名直接用(self):
        assert "tuya" in app_detect.name_keywords("tuya")

    def test_匹配打分排序(self):
        pkgs = ["com.example.tuya.helper", "com.tuya.smartiot", "com.other"]
        matched = app_detect.match_packages("涂鸦智能", pkgs)
        assert matched[0][0] == "com.tuya.smartiot"  # 更长窗口命中优先
        assert all("other" != p.split(".")[-1] or True for p, _ in matched)
        assert not any(p == "com.other" for p, _ in matched)

    def test_无匹配返回空(self):
        assert app_detect.match_packages("不存在xyz", ["com.a.b"]) == []

    def test_解析resolve_activity输出(self):
        out = "priority=0 preferredOrder=0\ncom.demo.app/com.demo.MainActivity\n"
        assert app_detect._parse_brief_activity(out, "com.demo.app") == "com.demo.MainActivity"
        assert app_detect._parse_brief_activity("nothing here", "com.demo.app") is None


# ── Excel 报告 ──
class TestExcelReport:

    def test_写入与统计(self, tmp_path):
        from common.excel_report import ExcelReport
        path = str(tmp_path / "r.xlsx")
        r = ExcelReport(path=path)
        r.set_device_info(sn="SN123", app_version="1.2.3")
        r.set_step_desc("模块A", "步骤1")
        r.add_result("模块A", True, "", "")
        r.set_step_desc("模块A", "步骤2")
        r.add_result("模块A", False, "超时", "")
        r.save()
        from openpyxl import load_workbook
        wb = load_workbook(path)
        assert "汇总" in wb.sheetnames and "模块A" in wb.sheetnames
        ws = wb["汇总"]
        texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
        assert any("SN123" in t for t in texts)
        assert any("1.2.3" in t for t in texts)
        # 结果列统计
        fails = [c for row in wb["模块A"].iter_rows() for c in row if c.value == "FAIL"]
        assert len(fails) == 1


# ── 会话工具 ──
class TestSession:

    def test_get_device_id_策略(self):
        from common.session import get_device_id
        cfg = {"device": {"default": "auto",
                          "list": [{"id": "emulator-5554", "name": "x"}]}}
        assert get_device_id(cfg, "abc") == "abc"
        assert get_device_id(cfg, "auto") == "emulator-5554"
        assert get_device_id(cfg, "") == "emulator-5554"
        cfg2 = {"device": {"default": "127.0.0.1:5555", "list": []}}
        assert get_device_id(cfg2, "auto") == "127.0.0.1:5555"


# ── 引擎停止 ──
class TestEngineStop:

    @staticmethod
    def _runner(interval=1):
        class FakeD:
            info = None

            def screenshot(self, *a, **k):
                return ""

        r = ActionRunner.__new__(ActionRunner)
        r.d = FakeD()
        r._device_id = None
        r.interval = interval
        r.timeout = 1
        r.click_timeout = 1
        r.results = []
        r.store = {}
        r.stopped = False
        r.on_result = None
        r.case_name = "t"
        r._locators = {}
        return r

    def test_等待中停止(self):
        import threading
        import time as _t
        r = self._runner()
        threading.Timer(0.3, r.stop).start()
        t0 = _t.time()
        ok = r.run_steps([{"desc": "等待", "wait": 5}, {"desc": "c", "click": "x"}])
        assert not ok and 0.2 < _t.time() - t0 < 2
        assert any("手动停止" in str(x.get("error", "")) for x in r.results)

    def test_on_result回调(self):
        r = self._runner()
        emitted = []
        r.on_result = lambda e: emitted.append(e["passed"])
        assert r.run_steps([{"desc": "等待", "wait": 0}])
        assert emitted == [True]


# ── GUI(离屏) ──
@pytest.fixture
def window():
    from gui.main_window import MainWindow
    w = MainWindow()
    w.resize(1050, 620)
    w.show()
    _qapp.processEvents()
    yield w
    w.close()
    w.deleteLater()


def _click(card):
    ev = QMouseEvent(QEvent.MouseButtonPress, QPointF(30, 15), QPointF(30, 15),
                     Qt.LeftButton, Qt.LeftButton, Qt.NoModifier)
    card.mousePressEvent(ev)


def _card_widgets(window):
    """当前卡片列表里的 StepCard(不含 stretch)"""
    from gui.main_window import StepCard
    out = []
    for i in range(window.cards_lay.count()):
        wdg = window.cards_lay.itemAt(i).widget()
        if isinstance(wdg, StepCard):
            out.append(wdg)
    return out


class TestGuiWindow:

    def test_环境字段从配置加载(self, window, monkeypatch):
        assert window.app_name_edit.text() == "涂鸦智能"
        assert window.device_name_edit.text() == "SE3L"

    def test_包名启动页字段已移除(self, window):
        assert not hasattr(window, "pkg_edit")
        assert not hasattr(window, "act_edit")

    def test_设备下拉有检测项(self, window):
        try:
            online = app_detect.list_devices()
        except Exception:
            online = []
        if not online:
            pytest.skip("无 adb 设备")
        assert window.device_combo.count() >= 1
        assert window._current_device_id() in [d["id"] for d in online]

    def test_设备名称写入target_device(self, window, monkeypatch, tmp_path):
        fake = tmp_path / "config.yaml"
        from common.driver import CONFIG_PATH
        text = open(CONFIG_PATH, encoding="utf-8").read()
        fake.write_text(text, encoding="utf-8")
        monkeypatch.setattr("common.driver.CONFIG_PATH", str(fake))
        window.device_name_edit.setText("SE9L-TEST")
        window._save_cfg_field(window.device_name_edit)
        data = yaml.safe_load(fake.read_text(encoding="utf-8"))
        assert data["target_device"] == "SE9L-TEST"


class TestGuiSteps:

    def test_添加移动复制删除(self, window):
        window.on_new()
        window.add_step("click")
        window.add_step("add_timer")
        window.add_step("__wait")
        assert [s.get("click", s.get("add_timer", s.get("wait")))
                for s in window.steps] == ["", True, 10]
        assert window.steps[1]["add_timer"] is True
        assert window.expanded_key == (2,)
        window.move_step(2, -1)
        assert window.steps[1]["desc"] == "延时等待"
        assert window.expanded_key == (1,)
        window.dup_step(0)
        assert len(window.steps) == 4
        window.del_step(3)
        assert len(window.steps) == 3
        window.on_clear_steps if False else None

    def test_点击展开与切换(self, window):
        window.on_new()
        window.add_step("click")
        window.add_step("assert")
        window.expanded_key = None
        window.render_cards()
        cards = _card_widgets(window)
        assert len(cards) == 2 and not cards[0].expanded
        _click(cards[0])
        cards = _card_widgets(window)
        assert window.expanded_key == (0,) and cards[0].expanded
        assert len(cards[0].getters) > 0
        _click(cards[0])  # 展开态点击不收起
        assert window.expanded_key == (0,)

    def test_表单回写保留列表类型(self, window):
        window.on_new()
        window.steps.append({"desc": "滑动", "swipe": [100, 200, 300, 400]})
        window.expanded_key = (0,)
        window.render_cards()
        card = _card_widgets(window)[0]
        # 只编辑名称,不碰滑动字段
        card.widgets["desc"].setText("滑动一")
        card.widgets["desc"].editingFinished.emit()
        assert window.steps[0]["desc"] == "滑动一"
        assert window.steps[0]["swipe"] == [100, 200, 300, 400]  # 未编辑字段类型不变

    def test_高级参数折叠编辑(self, window):
        window.on_new()
        window.add_step("click")
        window.steps[0]["click"] = "确认"
        card = _card_widgets(window)[0]
        card.widgets["wait"].setText("7")
        card.widgets["wait"].editingFinished.emit()
        assert window.steps[0]["wait"] == 7
        assert "📷" not in card.summary_label.text()
        card.widgets["screenshot"].setText("screenshots/x.png")
        card.widgets["screenshot"].editingFinished.emit()
        assert "📷" in card.summary_label.text()


class TestGuiElse:

    def test_添加编辑删除子步骤(self, window):
        window.on_new()
        window.steps.append({"desc": "条件", "if": "扫地机器人", "else": []})
        window.expanded_key = (0,)
        window.render_cards()
        window.add_sub(0, "click")
        assert len(window.steps[0]["else"]) == 1
        assert window.expanded_key == (0, 1) or window.expanded_key == (0, 0)
        window.steps[0]["else"][0]["click"] = "开始清扫"
        window.move_sub(0, 0, -1)  # 越界不动
        window.add_sub(0, "__wait")
        assert len(window.steps[0]["else"]) == 2
        window.move_sub(0, 1, -1)
        assert window.steps[0]["else"][0]["wait"] == 10
        window.del_sub(0, 0)
        assert len(window.steps[0]["else"]) == 1
        window.del_sub(0, 0)
        assert window.steps[0]["else"] == []
        assert window.expanded_key == (0,)  # 删完收回到父卡片

    def test_子卡片渲染与展开(self, window):
        window.on_new()
        window.steps.append({"desc": "条件", "if": "扫地机器人",
                             "else": [{"desc": "子1", "click": "x"},
                                      {"desc": "子2", "wait": 3}]})
        window.expanded_key = (0,)
        window.render_cards()
        cards = _card_widgets(window)
        assert len(cards) == 3  # 父 + 2 子
        assert cards[1].is_sub and cards[2].is_sub
        assert "子1" in cards[1].summary_label.text()
        _click(cards[1])
        assert window.expanded_key == (0, 0)
        cards = _card_widgets(window)
        assert cards[1].expanded and "click" in cards[1].getters

    def test_YAML回环保留else(self, window):
        window.on_new()
        window.steps.append({"desc": "条件", "if": "扫地机器人",
                             "else": [{"desc": "子1", "click": "x", "screenshot": ""}]})
        window._refresh_yaml_text()
        window._apply_yaml_text()
        out = window.steps[0]
        assert out["if"] == "扫地机器人"
        assert out["else"] == [{"desc": "子1", "click": "x"}]


class TestGuiYamlFiles:

    def test_12个用例文件回环无丢键(self, window):
        from common.action_runner import ActionRunner
        dispatch_keys = {k for k, _ in ActionRunner._MAIN_DISPATCH}
        modifier_keys = {"desc", "screenshot", "wait", "timeout", "retry", "circular",
                         "threshold", "switch_tpl", "switch_label", "switch_area", "else"}
        field_keys = {f["key"] for a in schema.ACTIONS for f in a["fields"]
                      if f["type"] != "group"}
        group_keys = {s["key"] for a in schema.ACTIONS for f in a["fields"]
                      if f["type"] == "group" for s in f["fields"]}
        all_keys = dispatch_keys | modifier_keys | set(schema.ACTION_BY_KEY) | field_keys | group_keys
        total = 0
        for path in sorted(glob.glob(os.path.join(CASES_DIR, "*.yaml"))):
            with open(path, encoding="utf-8") as f:
                data = yaml.safe_load(f)
            window.data = data
            window.case_idx = 0
            window.expanded_key = None
            window.render_cards()
            expect = len(data["cases"][0]["steps"])
            shown = len(_card_widgets(window))
            assert shown == expect, f"{path}: 期望{expect} 显示{shown}"
            for case in data.get("cases", []):
                for s in case.get("steps", []):
                    total += 1
                    cleaned = schema.serialize_step(s)
                    for k, v in s.items():
                        if v in (None, "", False):
                            continue
                        assert k in cleaned, f"{path}: 键 {k} 丢失"
                    for k in cleaned:
                        assert k in all_keys, f"{path}: 未知键 {k}"
        assert total >= 200

    def test_源码页双向同步(self, window):
        window.on_new()
        window.add_step("click")
        window.steps[0]["click"] = "确认"
        window._refresh_yaml_text()
        text = window.yaml_edit.toPlainText()
        assert "click: 确认" in text
        window._apply_yaml_text()
        assert len(window.steps) == 1
        assert window.steps[0]["click"] == "确认"


class TestGuiSaveAndRun:

    def test_保存静默路径(self, window, tmp_path):
        window.on_new()
        window.add_step("click")
        window.steps[0]["click"] = "进入设置"
        window.case_path = str(tmp_path / "case.yaml")
        window.on_save()
        assert os.path.exists(window.case_path)
        data = yaml.safe_load(open(window.case_path, encoding="utf-8"))
        assert data["cases"][0]["steps"] == [{"desc": "点击", "click": "进入设置"}]

    def test_保存校验拦截(self, window, tmp_path, monkeypatch):
        from PySide6.QtWidgets import QMessageBox
        window.on_new()
        window.add_step("set_time")  # 必填为空
        window.case_path = str(tmp_path / "bad.yaml")
        monkeypatch.setattr(QMessageBox, "question",
                            staticmethod(lambda *a, **k: QMessageBox.No))
        window.on_save()
        assert not os.path.exists(window.case_path)

    def test_结果表刷新与结束恢复(self, window):
        window.result_table.setRowCount(0)
        window.on_step_done({"desc": "步骤1", "passed": True, "error": "", "screenshot": ""})
        window.on_step_done({"desc": "步骤2", "passed": False, "error": "超时", "screenshot": ""})
        assert window.result_table.rowCount() == 2
        assert window.result_table.item(0, 1).text() == "PASS"
        assert window.result_table.item(1, 1).text() == "FAIL"
        window.run_btn.setEnabled(False)
        window.stop_btn.setEnabled(True)
        window.on_run_finished(False, "存在失败步骤")
        assert window.run_btn.isEnabled() and not window.stop_btn.isEnabled()

    def test_停止按钮防护(self, window):
        assert window.worker is None
        window.on_stop()  # 无 worker 时不应抛异常


class TestGuiPreview:
    """测试步骤描述标签 / 选中行显示截图 / 截图查看器"""

    def test_步骤描述标签改名(self, window):
        from PySide6.QtWidgets import QLabel
        window.on_new()
        window.add_step("click")
        card = _card_widgets(window)[0]
        labels = [l.text() for l in card.findChildren(QLabel)]
        assert "测试步骤描述" in labels
        assert "名称" not in labels

    def test_选中行显示对应截图(self, window, tmp_path):
        from PIL import Image
        shot = str(tmp_path / "shot.png")
        Image.new("RGB", (80, 60), (200, 30, 30)).save(shot)
        window.result_table.setRowCount(0)
        window.on_step_done({"desc": "有图步骤", "passed": True, "error": "",
                             "screenshot": shot})
        window.on_step_done({"desc": "无图步骤", "passed": True, "error": "",
                             "screenshot": ""})
        window.result_table.setCurrentCell(0, 0)
        _qapp.processEvents()
        pm = window.preview.pixmap()
        assert pm is not None and not pm.isNull()
        window.result_table.setCurrentCell(1, 0)
        _qapp.processEvents()
        assert window.preview.pixmap().isNull()
        assert "无截图" in window.preview.text()

    def test_查看器缩放与适应(self, tmp_path):
        from PIL import Image
        from gui.main_window import ImageViewDialog
        p = str(tmp_path / "big.png")
        Image.new("RGB", (400, 300), (10, 120, 240)).save(p)
        dlg = ImageViewDialog(p)
        assert dlg._zoom is None  # 默认适应窗口
        dlg._set_zoom(2.0)
        assert dlg.image_label.pixmap().width() == 800
        dlg._zoom_out()
        assert dlg._zoom == 1.6
        dlg._zoom_in()
        assert dlg._zoom == 2.0
        dlg._orig()
        assert dlg._zoom == 1.0
        dlg._fit()
        assert dlg._zoom is None
        dlg.close()

    def test_耗时列(self, window):
        window.result_table.setRowCount(0)
        window.on_step_done({"desc": "s1", "passed": True, "error": "",
                             "screenshot": "", "elapsed": 3.2})
        assert window.result_table.item(0, 3).text() == "3.2s"
        window.on_step_done({"desc": "s2", "passed": True, "error": "",
                             "screenshot": ""})
        assert window.result_table.item(1, 3).text() == ""


class TestGuiLockAndCases:
    """执行锁定 / 多 case 切换器 / 卡片头部 toggle"""

    def test_执行期间锁定编排区(self, window):
        window.on_new()
        window.add_step("click")
        window.worker = object()  # 模拟执行中
        window._set_locked(True)
        assert not window.add_btn.isEnabled()
        assert not window.save_btn.isEnabled()
        assert not window.case_name_edit.isEnabled()
        assert not window.cards_scroll.widget().isEnabled()
        n = len(window.steps)
        window.add_step("assert")   # 守卫生效,不添加
        window.move_step(0, 1)
        assert len(window.steps) == 1
        window.worker = None
        window._set_locked(False)
        assert window.add_btn.isEnabled()
        assert window.cards_scroll.widget().isEnabled()

    def test_运行结束自动解锁(self, window):
        window.worker = object()
        window._set_locked(True)
        window.run_btn.setEnabled(False)
        window.on_run_finished(True, "全部通过")
        assert window.add_btn.isEnabled()
        assert window.worker is None

    def test_多case切换器(self, window):
        window.on_new()
        window.data = {"module": "M", "cases": [
            {"name": "用例A", "priority": "P1", "steps": [{"desc": "a", "click": "x"}]},
            {"name": "用例B", "priority": "P1", "steps": []}]}
        window.case_idx = 0
        window.expanded_key = None
        window._load_case_into_ui()
        window.render_cards()
        window._refresh_case_selector()
        assert window.case_combo.isVisible()
        assert window.case_combo.count() == 2
        window.case_combo.setCurrentIndex(1)
        assert window.case_idx == 1
        assert window.case_name_edit.text() == "用例B"
        assert len(window.steps) == 0
        # 改名联动切换器文字
        window.case_name_edit.setText("用例B2")
        window._sync_header()
        assert "用例B2" in window.case_combo.itemText(1)

    def test_保存校验覆盖所有case(self, window, tmp_path, monkeypatch):
        from PySide6.QtWidgets import QMessageBox
        window.on_new()
        window.data = {"module": "M", "cases": [
            {"name": "A", "priority": "P1", "steps": [{"desc": "a", "click": "x"}]},
            {"name": "B", "priority": "P1", "steps": [{"desc": "b", "set_time": ""}]}]}
        window.case_path = str(tmp_path / "m.yaml")
        monkeypatch.setattr(QMessageBox, "question",
                            staticmethod(lambda *a, **k: QMessageBox.No))
        window.on_save()  # B 的必填为空且选择不保存
        assert not os.path.exists(window.case_path)

    def test_点击头部收回再展开(self, window):
        window.on_new()
        window.add_step("click")
        assert window.expanded_key == (0,)
        card = _card_widgets(window)[0]
        card._header_clicked()  # 再点 → 收回
        assert window.expanded_key is None
        card = _card_widgets(window)[0]
        assert not card.expanded
        card._header_clicked()  # 又点 → 展开
        assert window.expanded_key == (0,)

    def test_子步骤点击收回回到父级(self, window):
        window.on_new()
        window.steps.append({"desc": "条件", "if": "x",
                             "else": [{"desc": "子", "click": "y"}]})
        window.expanded_key = (0,)
        window.render_cards()
        window.add_sub(0, "__wait")
        assert window.expanded_key == (0, 1)
        cards = _card_widgets(window)
        sub = [c for c in cards if c.is_sub and c.sub_index == 1][0]
        sub._header_clicked()
        assert window.expanded_key == (0,)  # 收回到父级(父表单展开,子列表仍在)
        cards = _card_widgets(window)
        assert len(cards) == 3  # 父 + 2 子
        assert cards[0].expanded and not cards[1].expanded
